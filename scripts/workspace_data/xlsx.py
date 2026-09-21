"""Read workbook cells with exact locations; never calculate formulas or follow links."""
from datetime import date, datetime, time, timedelta
import io
import math
import zipfile
from xml.etree.ElementTree import ParseError


def value_json(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return {'duration_seconds': value.total_seconds()}
    if type(value) is float and not math.isfinite(value):
        raise ValueError('Workbook contains a non-finite number')
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    raise ValueError('Workbook contains an unsupported cell value')


def extract_xlsx(content, max_records):
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError('Workbook extraction requires the recording extra with openpyxl 3.1.5') from exc
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 5000 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
                raise ValueError('Workbook exceeds expanded-size limits')
            if len({i.filename for i in infos}) != len(infos):
                raise ValueError('Workbook contains duplicate archive parts')
            for info in infos:
                if info.filename.endswith(('.xml', '.rels')):
                    raw = archive.read(info)
                    if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                        raise ValueError('Workbook XML entities are unsupported')
            media = sum(i.filename.startswith('xl/media/') for i in infos)
            charts = sum(i.filename.startswith('xl/charts/chart') and i.filename.endswith('.xml') for i in infos)
            external = sum(i.filename.startswith('xl/externalLinks/externalLink') and i.filename.endswith('.xml') for i in infos)
        source = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            cached = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        except Exception:
            source.close()
            raise
    except (zipfile.BadZipFile, KeyError, OSError, ParseError, TypeError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise ValueError('Could not read the Excel workbook') from exc
    records, sheets = [], []
    formulas = missing_cache = scanned_rows = scanned_cells = 0
    try:
        if len(source.worksheets) > 100:
            raise ValueError('Workbook exceeds 100 worksheet limit')
        for sheet in source.worksheets:
            values = cached[sheet.title]
            sheet.reset_dimensions(); values.reset_dimensions()
            count = 0
            for row, value_row in zip(sheet.iter_rows(), values.iter_rows()):
                scanned_rows += 1
                scanned_cells += len(row)
                if scanned_cells > 250000:
                    raise ValueError('Workbook exceeds 250,000 scanned cell positions')
                if scanned_rows > 100000:
                    raise ValueError('Workbook exceeds 100,000 scanned rows')
                for cell, cached_cell in zip(row, value_row):
                    if cell.value is None:
                        continue
                    is_formula = cell.data_type == 'f'
                    if is_formula and not isinstance(cell.value, str):
                        raise ValueError('Array or data-table formulas require a values export for this intake version')
                    value = value_json(cached_cell.value if is_formula else cell.value)
                    formulas += int(is_formula)
                    missing_cache += int(is_formula and value is None)
                    records.append({'locator': {'sheet': sheet.title, 'cell': cell.coordinate, 'row': cell.row, 'column': cell.column_letter, 'sheet_state': sheet.sheet_state},
                                    'data': {'sheet': sheet.title, 'cell': cell.coordinate, 'value': value,
                                             'formula': cell.value if is_formula else None,
                                             'value_origin': ('cached_formula_result_unverified' if value is not None else 'formula_result_unavailable') if is_formula else 'stored_cell',
                                             'cell_type': cell.data_type, 'number_format': cell.number_format}})
                    count += 1
                    if len(records) > max_records:
                        raise ValueError('Workbook exceeds cell record limits')
            sheets.append({'name': sheet.title, 'state': sheet.sheet_state, 'cells': count})
    except (ParseError, KeyError, TypeError) as exc:
        raise ValueError('Could not read workbook cells') from exc
    finally:
        source.close(); cached.close()
    if not records:
        raise ValueError('Workbook has no readable cell values or formulas')
    return {'kind': 'workbook', 'status': 'extracted', 'extractor': f'openpyxl-{openpyxl.__version__}-cells-v1', 'requires_structuring': True,
            'records': records, 'workbook_coverage': {'sheets': sheets, 'cells': len(records), 'formulas': formulas,
                'formulas_without_cached_value': missing_cache, 'date_system': str(source.epoch.date()),
                'embedded_media_not_read': media, 'charts_not_read': charts, 'external_links_not_followed': external,
                'limitation': 'Nonempty cells and formula text from every worksheet, including hidden sheets. Dates use ISO values. Formula results are saved caches, may be stale, and are never recalculated. Blank cells, merged layout, charts, images, comments, formatting semantics, and external linked data are not interpreted.'}}
